# -*- coding: utf-8 -*-
"""Convertor Markdown -> DOCX pentru raportul de sustenabilitate.

Suportă: titluri (#/##/###), paragrafe, liste (-, 1.), tabele pipe, imagini,
blocuri de cod, **bold**, *italic*, `code`, [text](url), referințe [S1]
(legate către URL-ul din definițiile de subsol).
Parsarea este comună cu md2pdf.py (vezi md_comun.py).
"""
import os
import sys
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from md_comun import (citeste, definitii_referinte, tokeni_inline, blocuri,
                      titlu, imagine, randuri_tabel, element_lista, url_public)

MD_PATH = sys.argv[1]
DOCX_PATH = sys.argv[2]
BASE = os.path.dirname(os.path.abspath(MD_PATH))

text = citeste(MD_PATH)
refdefs = definitii_referinte(text)

doc = Document()
stil = doc.styles["Normal"]
stil.font.name = "Calibri"
stil.font.size = Pt(10.5)
stil.paragraph_format.space_after = Pt(6)
for sec in doc.sections:
    sec.left_margin = Cm(2.0); sec.right_margin = Cm(2.0)
    sec.top_margin = Cm(1.8); sec.bottom_margin = Cm(1.8)
    sec.page_width = Cm(21.0); sec.page_height = Cm(29.7)  # A4


def add_hyperlink(par, label, url):
    part = par.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    color = OxmlElement("w:color"); color.set(qn("w:val"), "0563C1"); rPr.append(color)
    u = OxmlElement("w:u"); u.set(qn("w:val"), "single"); rPr.append(u)
    new_run.append(rPr)
    t = OxmlElement("w:t"); t.text = label
    new_run.append(t)
    h.append(new_run)
    par._p.append(h)


def add_inline(par, s):
    for tok in tokeni_inline(s, refdefs):
        tip = tok[0]
        if tip == "text":
            par.add_run(tok[1])
        elif tip == "bold":
            r = par.add_run(tok[1]); r.bold = True
        elif tip == "italic":
            r = par.add_run(tok[1]); r.italic = True
        elif tip == "code":
            r = par.add_run(tok[1]); r.font.name = "Consolas"; r.font.size = Pt(9.5)
        elif tip in ("link", "ref"):
            url = url_public(tok[2])
            if url:
                add_hyperlink(par, tok[1], url)
            else:
                par.add_run(tok[1])


def add_image(par, path):
    p = os.path.normpath(os.path.join(BASE, path))
    if os.path.exists(p):
        try:
            par.add_run().add_picture(p, width=Cm(16.0))
        except Exception as e:
            par.add_run(f"[imagine: {path} — {e}]")
    else:
        par.add_run(f"[imagine lipsă: {path}]")
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER


first_heading = True
for tip, payload in blocuri(text):
    if tip == "cod":
        par = doc.add_paragraph()
        r = par.add_run(payload)
        r.font.name = "Courier New"; r.font.size = Pt(9.5)
        par.paragraph_format.left_indent = Cm(0.6)
    elif tip == "par":
        par = doc.add_paragraph()
        add_inline(par, payload)
    elif tip == "head":
        nivel, titlu_txt = titlu(payload)
        if nivel == 1 and first_heading:
            doc.add_heading(titlu_txt, level=0)
            first_heading = False
        else:
            t = doc.add_heading("", level=min(nivel, 4))
            add_inline(t, titlu_txt)
    elif tip == "img":
        img = imagine(payload)
        if img:
            alt, cale = img
            par = doc.add_paragraph()
            add_image(par, cale)
            cap = doc.add_paragraph(alt)
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap.runs[0].italic = True
            cap.runs[0].font.size = Pt(9)
    elif tip == "table":
        celule = randuri_tabel(payload)
        if celule:
            tbl = doc.add_table(rows=len(celule), cols=len(celule[0]))
            tbl.style = "Table Grid"
            for ri, row in enumerate(celule):
                for ci, txt in enumerate(row):
                    cell = tbl.cell(ri, ci)
                    cell.text = ""
                    par = cell.paragraphs[0]
                    par.paragraph_format.space_after = Pt(2)
                    add_inline(par, txt)
                    for r in par.runs:
                        r.font.size = Pt(9)
                        if ri == 0:
                            r.bold = True
            doc.add_paragraph()
    elif tip in ("bullet", "numar"):
        prefix, continut = element_lista(tip, payload)
        if tip == "bullet":
            par = doc.add_paragraph(style="List Bullet")
        else:
            par = doc.add_paragraph()
            par.paragraph_format.left_indent = Cm(0.9)
            par.paragraph_format.first_line_indent = Cm(-0.5)
        add_inline(par, prefix + continut)
        if tip == "numar" and par.runs:
            par.runs[0].bold = True

doc.save(DOCX_PATH)
print("Scris:", DOCX_PATH)
print("Referințe găsite:", len(refdefs))
